"""Build the per-driver attribution record and render the PRISM Analysis texts from it.

    python3 scripts/build_driver_analysis.py                  record, parity fixture and texts
    python3 scripts/build_driver_analysis.py --reference      also re-run the 3 x 50,000 reference run
    python3 scripts/build_driver_analysis.py --check          CI: the committed files match this code
    python3 scripts/build_driver_analysis.py --workbook FILE   record which production run is reproduced

Files
  data/drivers/driver_attribution.json   the figures (generated)
  data/drivers/driver_analysis.json      one hand-written template per driver and the text
                                         rendered from it (generated)
  tests/frontend/fixtures/driverAttribution.fixture.json
                                         parity fixture for lib/driverAttribution.ts, the
                                         port of the decomposition the category drill-down uses

The files are built on the default cell weights, the estimated gross-profit
mix, and the templates are written for it: several name the estimate and
compare shares of it in words. There is deliberately no option to render them
on other shares; readings on the finance shares need templates written for
those shares first (docs/DRIVER_ANALYSIS.md, "On the finance shares").

What the record is
  The engine's own arithmetic, evaluated at the score means instead of at
  Monte Carlo draws, and taken apart driver by driver. For driver j, cell
  (category c, region r) and path year y the additive term of the
  compounding factor is

      a[j, c, r, y] = force_weight[f] x attenuation[f] x dampening[f, c, r]
                      x E[p_j] x gp1_j x sign_j x E[materialisation_j(y)]
                      x category_exposure[j, c] / 5 x regional_exposure[j, r] / 5

  with E[p_j] = probability / 6 (the Beta-prior mean) and the
  materialisation averaged over the driver's peak-year jitter grid. This is
  the quantity the engine sums per force in `_compute_all_paths_vectorized`.
  The force sums compound multiplicatively, so the cell shift at the means
  is prod_f (1 + sum_j a) - 1, cross-terms included; the cross-terms (a few
  percent of the cell shift) are spread over the drivers of the cell in
  proportion to their additive terms, so the drivers of a cell add up to the
  cell shift exactly. The roll-up is the engine's: portfolio = sum over
  cells of the gross-profit share times the cell shift, category = the
  share-weighted mean of the category's row.

  The sum at the means is the run's terminal-year mean to a few hundredths
  of a percent; the residual is the covariance between the force factors,
  which the Monte Carlo averages and an evaluation at the means does not.
  The dashboard's headline is the median, and the portfolio distribution is
  skewed, so every contribution is scaled by one factor, the reference run's
  terminal-year median over the sum at the means, and the contributions add
  up to the headline the reader sees. The factor is the skew of the
  distribution, not an approximation error; the record keeps it and the
  reference run it comes from.

  Each driver's term is linear in its probability draw, so the mean of the
  Monte Carlo term equals the term at the mean: this is an exact
  expected-value decomposition of the additive part, not a fitted
  attribution, and it is labelled attribution wherever it is shown.

  Per driver the record also carries where the contribution lands (cells,
  categories, regions); when it lands (the share of the terminal-year
  contribution present in each path year, which is the materialisation
  ratio because the additive term factorises); the driver's own P10 to P90
  range at the horizon and its share of the portfolio band (200,000 draws of
  the engine's own copula and jitter samplers, Euler allocation over the
  additive sum, see `band_analysis`); the pass-through it faces, the overlap
  correction in its cells and the same-force drivers it shares them with;
  its reach, the share of the gross-profit mix its exposures cover (sum over
  cells of cell share x category exposure / 5 x regional exposure / 5); and,
  over all 52 cells, the cell where it is the largest share of the cell's
  own shift.

Why a generated record: the texts quote these figures, and a quoted figure
must never drift from the model. `--check` recomputes the record from the
driver base and the configuration, re-renders every text and fails when a
score, an exposure, the overlap correction or the cell weights changed
without the analysis being regenerated. It also checks that the reference
run was made on the current inputs (their digest) and that it reproduces
the production run recorded with --workbook.

After a re-scoring: run scripts/run_production.py, then this script with
--reference --workbook FILE, review the re-rendered texts, run
scripts/generate_seed.py and scripts/sync_analysis.py.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import json
import os
import sys

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, REPO)

# The record is a function of the driver base and the configuration, never of
# a database: blank the URLs before prism.env_loader can read .env.
os.environ["DATABASE_URL"] = ""
os.environ["POSTGRES_URL"] = ""

import numpy as np  # noqa: E402

from prism.config import (DEFAULT_PATH_YEARS, FORCES, PRODUCTION_CHAINS, PRODUCTION_ITERATIONS,  # noqa: E402
                          PRODUCTION_SEED, REGIONS, ModelConfig, display_code)
from prism.drivers import DriverSet  # noqa: E402
from prism.simulation.engine import MonteCarloEngine, NUMERICS_BACKEND  # noqa: E402
from scripts.driver_analysis_render import render_all  # noqa: E402

DATA_DIR = os.path.join(REPO, "data", "drivers")
ATTRIBUTION_PATH = os.path.join(DATA_DIR, "driver_attribution.json")
ANALYSIS_PATH = os.path.join(DATA_DIR, "driver_analysis.json")
#: Parity fixture for lib/driverAttribution.ts (the category drill-down uses
#: the same decomposition): the inputs and the unscaled expected outputs.
FIXTURE_PATH = os.path.join(REPO, "tests", "frontend", "fixtures", "driverAttribution.fixture.json")

TERMINAL_YEAR = DEFAULT_PATH_YEARS[-1]
#: The draws behind each driver's band share, with a fixed seed so the shares
#: are reproducible (any fixed value would do; changing it moves the shares).
BAND_DRAWS = 200_000
BAND_SEED = 20260917
#: How many of its largest cells, and same-force partners per cell, the
#: record keeps for each driver.
TOP_CELLS = 6
TOP_PARTNERS = 3


# ── drivers ─────────────────────────────────────────────────────────────

def load_drivers() -> list:
    """The drivers of the base, ordered by id as the production run loads them."""
    from scripts.generate_seed import base_drivers
    return base_drivers()


# ── the decomposition ──────────────────────────────────────────────────

def decompose(config: ModelConfig, drivers: list) -> dict:
    """Evaluate the engine's arithmetic at the score means, per driver."""
    eng = MonteCarloEngine(config, seed=PRODUCTION_SEED)
    cats = list(config.category_names)
    years = [int(y) for y in config.path_years]
    n, nc, nr, ny = len(drivers), len(cats), len(REGIONS), len(years)
    ncell = nc * nr
    forces = list(FORCES)
    nf = len(forces)

    fw = np.array([config.force_weights.get(f, 1.0 / nf) for f in forces])
    att = np.array([eng.effective_attenuation[f] for f in forces])
    fidx = np.array([forces.index(drv.force) for drv in drivers])

    # Composite cell exposure, exactly as the engine builds it.
    region_w = np.zeros((n, nr))
    regionless = []
    for j, drv in enumerate(drivers):
        reg = drv.regional_exposure or {}
        vals = np.array([min(float(reg.get(r, 0.0)), 5.0) / 5.0 for r in REGIONS])
        if vals.sum() <= 0:
            vals = np.ones(nr)
            regionless.append(drv.id)
        region_w[j, :] = vals
    E = np.zeros((n, ncell))
    for j, drv in enumerate(drivers):
        for c, cat in enumerate(cats):
            exp = drv.category_exposure.get(cat, 0)
            ce = min(exp, 5) / 5.0 if exp > 0 else 0.0
            if ce <= 0:
                continue
            for r in range(nr):
                E[j, c * nr + r] = ce * region_w[j, r]

    # Magnitude-weighted within-force dampening, exactly as the engine.
    abs_score = np.array([abs(float(drv.normalized_score)) for drv in drivers])
    contrib = abs_score[:, None] * E
    damp = np.ones((nf, ncell))
    n_eff_all = np.zeros((nf, ncell))
    wf_overlap = config.within_force_overlap
    for f, force in enumerate(forces):
        mask = fidx == f
        if not mask.any():
            continue
        ov = wf_overlap.get(force, 0.0)
        M = contrib[mask, :]
        s1 = M.sum(axis=0)
        s2 = (M * M).sum(axis=0)
        n_eff = np.where(s2 > 0.0, (s1 * s1) / np.where(s2 > 0.0, s2, 1.0), 0.0)
        n_eff_all[f, :] = n_eff
        if ov <= 0:
            continue
        n_eff_safe = np.maximum(n_eff, 1.0)
        d = np.where(n_eff > 1.0, 1.0 - ov * (n_eff - 1.0) / n_eff_safe, 1.0)
        damp[f, :] = np.clip(d, 0.1, 1.0)

    # The schedule averaged over each driver's peak-jitter grid.
    sched, offsets, off_probs = eng.peak_schedule_table(drivers)     # (n, n_off, ny)
    e_mat = np.einsum("jk,jky->jy", off_probs, sched)                # (n, ny)
    s_mean = np.array([float(drv.normalized_score) for drv in drivers])   # E[p] x gp1 x sign

    # Additive terms a[j, cell, y] and the per-cell compounding at the means.
    A = (fw[fidx][:, None, None] * att[fidx][:, None, None]
         * damp[fidx, :][:, :, None] * s_mean[:, None, None]
         * E[:, :, None] * e_mat[:, None, :])
    F = np.zeros((nf, ncell, ny))
    for f in range(nf):
        F[f] = A[fidx == f].sum(axis=0)
    # the engine floors every force factor at 0 (a factor cannot take a cell
    # below -100 percent); inert at the means on any realistic base
    cell_shift = np.prod(np.maximum(1.0 + F, 0.0), axis=0) - 1.0     # (ncell, ny)
    additive = F.sum(axis=0)
    spread = np.where(np.abs(additive) > 1e-15,
                      cell_shift / np.where(np.abs(additive) > 1e-15, additive, 1.0), 1.0)
    A_cell = A * spread[None, :, :]                                   # sums to cell_shift per cell

    W = eng.cell_weight_matrix()                                     # (nc, nr)
    Wn = eng.row_normalized_cell_weights()
    w_cell = W.reshape(-1)
    wn_cell = Wn.reshape(-1)

    port = (A_cell * w_cell[None, :, None]).sum(axis=1)               # (n, ny) portfolio pp (as fraction)
    cat_pp = np.zeros((n, nc, ny))
    for c in range(nc):
        sl = slice(c * nr, (c + 1) * nr)
        cat_pp[:, c, :] = (A_cell[:, sl, :] * wn_cell[sl][None, :, None]).sum(axis=1)
    reg_pp = np.zeros((n, nr, ny))
    for r in range(nr):
        idx = [c * nr + r for c in range(nc)]
        reg_pp[:, r, :] = (A_cell[:, idx, :] * w_cell[idx][None, :, None]).sum(axis=1)

    return {
        "engine": eng, "cats": cats, "years": years, "forces": forces,
        "fw": fw, "att": att, "fidx": fidx, "E": E, "damp": damp, "n_eff": n_eff_all,
        "sched": sched, "offsets": offsets, "off_probs": off_probs, "e_mat": e_mat,
        "A_cell": A_cell, "cell_shift": cell_shift, "W": W, "Wn": Wn,
        "port": port, "cat_pp": cat_pp, "reg_pp": reg_pp, "regionless": regionless,
        "s_mean": s_mean,
    }


def band_analysis(config: ModelConfig, drivers: list, dec: dict, factor: float) -> tuple:
    """Each driver's own terminal-year contribution at P10 / P50 / P90, and its
    share of the portfolio band.

    200,000 draws of the engine's own copula sampler and peak-jitter sampler
    (one dedicated seed), so the draws carry the configured dependence
    structure. Per draw and driver the terminal-year term is
    K_j x sample_j x materialisation_j, with K_j the driver's constant from
    the decomposition (pass-through, dampening, exposure, cell shares, the
    cross-term spread of each cell and the reconciliation factor). The band
    share is the Euler allocation Cov(term_j, sum of terms) / Var(sum of
    terms): the shares add up to one, and a driver's share is the part of
    the band's variance that moves with it, including what it shares with
    correlated drivers. The additive sum is used as the total (the
    compounding cross-terms are second order); the record also keeps the
    standard deviation of that sum next to the reference run's band so the
    approximation can be judged."""
    years = dec["years"]
    y_last = len(years) - 1
    A_cell, E, damp, fw, att, fidx, W = (dec["A_cell"], dec["E"], dec["damp"],
                                          dec["fw"], dec["att"], dec["fidx"], dec["W"])
    w_cell = W.reshape(-1)
    n = len(drivers)
    K = np.zeros(n)
    for j in range(n):
        f = fidx[j]
        term_const = fw[f] * att[f] * damp[f, :] * E[j, :]            # per cell, before score x materialisation
        a_raw = term_const * dec["s_mean"][j] * dec["e_mat"][j, y_last]
        with np.errstate(divide="ignore", invalid="ignore"):
            sp = np.where(np.abs(a_raw) > 1e-18,
                          A_cell[j, :, y_last] / np.where(np.abs(a_raw) > 1e-18, a_raw, 1.0), 1.0)
        K[j] = float((w_cell * term_const * sp).sum()) * factor

    eng = MonteCarloEngine(config, seed=BAND_SEED)
    R = eng.build_correlation_matrix(drivers)
    samples = eng.generate_copula_samples(drivers, R, BAND_DRAWS)      # p x gp1 x sign, (draws, n)
    off_idx = eng.draw_peak_offsets(dec["off_probs"], BAND_DRAWS)      # (draws, n)
    mat = dec["sched"][np.arange(n)[None, :], off_idx, y_last]          # (draws, n)
    terms = samples * mat * K[None, :]
    total = terms.sum(axis=1)
    var_total = float(np.var(total))
    out = []
    for j in range(n):
        q10, q50, q90 = np.percentile(terms[:, j], [10, 50, 90])
        cov = float(np.mean((terms[:, j] - terms[:, j].mean()) * (total - total.mean())))
        out.append({
            "p10_pp": round(float(q10) * 100, 4),
            "p50_pp": round(float(q50) * 100, 4),
            "p90_pp": round(float(q90) * 100, 4),
            "sigma_pp": round(float(np.std(terms[:, j])) * 100, 4),
            "band_share": round(cov / var_total, 5) if var_total > 0 else 0.0,
        })
    p10, p90 = np.percentile(total, [10, 90])
    return out, {"sigma_pp": round(float(np.sqrt(var_total)) * 100, 4),
                 "p10_pp": round(float(p10) * 100, 4), "p90_pp": round(float(p90) * 100, 4),
                 "draws": BAND_DRAWS, "seed": BAND_SEED}


def inputs_hash(config: ModelConfig, drivers: list) -> str:
    """A digest of everything the reference run depends on: the scoring
    state of every driver (its input-drift fingerprint, which includes the
    onset year), the force layers, the cell weights, the path and the engine
    version. The reference run is stamped with it, and a record whose inputs
    no longer match its reference run is refused, so that a re-scored base
    cannot be reconciled to a median it does not produce."""
    import hashlib
    from prism.input_drift import driver_fingerprint
    payload = {
        "drivers": driver_fingerprint(drivers),
        "force_weights": {f: round(float(config.force_weights[f]), 6) for f in FORCES},
        "per_force_attenuation": {f: round(float(config.per_force_attenuation[f]), 6) for f in FORCES},
        "within_force_overlap": {f: round(float(config.within_force_overlap.get(f, 0.0)), 6) for f in FORCES},
        "force_correlation_matrix": {a: {b: round(float(v), 6) for b, v in row.items()}
                                     for a, row in config.force_correlation_matrix.items()},
        "within_force_rho": round(float(config.within_force_rho), 6),
        "cell_weights": {c: {r: round(float(v), 8) for r, v in row.items()} for c, row in config.cell_weights.items()},
        "base_year": int(config.base_year), "path_years": list(config.path_years),
        "peak_year_jitter": int(config.peak_year_jitter),
        "model_version": MonteCarloEngine.MODEL_VERSION,
    }
    return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode("utf-8")).hexdigest()


def reference_run(config: ModelConfig, drivers: list, chains: int, iterations: int) -> dict:
    """The production recipe (prism.config PRODUCTION_*): run_multichain with the master seed."""
    db = DriverSet(drivers=list(drivers))
    eng = MonteCarloEngine(config, seed=PRODUCTION_SEED)
    res = eng.run_multichain(db, n_chains=chains, iterations=iterations)
    tot = res["totals"]["portfolio"]
    term = tot.get(TERMINAL_YEAR) or tot.get(str(TERMINAL_YEAR))
    cats = {}
    for cat, block in res["shift_matrix"].items():
        p = block["path"]
        cell = p.get(TERMINAL_YEAR) or p.get(str(TERMINAL_YEAR))
        cats[cat] = {k: round(float(cell[k]), 6) for k in ("p10", "median", "p90")}
    return {
        "seed": PRODUCTION_SEED, "chains": chains, "iterations": iterations,
        "pooled_iterations": chains * iterations,
        "numerics_backend": NUMERICS_BACKEND,
        "model_version": res["model_version"],
        "portfolio_terminal": {k: round(float(term[k]), 6) for k in ("p10", "p25", "median", "p75", "p90", "mean")},
        "categories_terminal": cats,
        "inputs_hash": inputs_hash(config, drivers),
        "run_at_utc": _dt.datetime.now(_dt.timezone.utc).isoformat(timespec="seconds"),
    }


def read_workbook(path: str) -> dict:
    """The production run as its QA workbook records it (prism/workbook.py).

    Run id, generation time, pooled iterations, model version and the 13
    terminal-year category quantiles. The record keeps them, so that `--check`
    proves the reference run reproduces the run the dashboard serves, and a
    reader of the record knows which run that is.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    meta = {}
    for row in wb["Metadata"].iter_rows(values_only=True):
        if row and row[0] is not None and len(row) > 1:
            meta[str(row[0]).strip()] = row[1]
    rows = list(wb["Shift Matrix"].iter_rows(values_only=True))
    header, sub = rows[0], rows[1]
    start = next(i for i, v in enumerate(header) if str(v) == str(TERMINAL_YEAR))
    assert tuple(str(x) for x in sub[start:start + 5]) == ("p10", "p25", "median", "p75", "p90"), sub[start:start + 5]
    cats = {}
    for r in rows[2:]:
        if r and r[0] is not None and r[start] is not None:
            p10, _, med, _, p90 = r[start:start + 5]
            cats[str(r[0])] = {"p10": round(float(p10), 6), "median": round(float(med), 6), "p90": round(float(p90), 6)}
    return {
        "id": int(meta["Run id"]) if meta.get("Run id") is not None else None,
        "generated": str(meta.get("Generated") or ""),
        "model_version": str(meta.get("Model version") or ""),
        "pooled_iterations": int(meta["Iterations"]) if meta.get("Iterations") is not None else None,
        "workbook": os.path.basename(path),
        "categories_terminal": cats,
    }


def reference_reproduces(reference: dict, persisted: dict, tol: float = 1.5e-6) -> list:
    """Names of the categories and quantiles where the reference run does not
    reproduce the persisted run (empty when it does)."""
    out = []
    for cat, q in persisted.get("categories_terminal", {}).items():
        ref = reference.get("categories_terminal", {}).get(cat)
        if ref is None:
            out.append(f"{cat}: not in the reference run"); continue
        for k in ("p10", "median", "p90"):
            if abs(float(ref[k]) - float(q[k])) > tol:
                out.append(f"{cat} {k}: reference {ref[k]} vs persisted {q[k]}")
    return out


# ── the record ─────────────────────────────────────────────────────────

def r4(x: float) -> float:
    return round(float(x), 4)


def build_record(config: ModelConfig, drivers: list, reference: dict, persisted: dict | None = None) -> dict:
    dec = decompose(config, drivers)
    cats, years, forces = dec["cats"], dec["years"], dec["forces"]
    y_last = len(years) - 1
    port_raw = dec["port"]
    det_total = float(port_raw[:, y_last].sum())
    median = float(reference["portfolio_terminal"]["median"])
    factor = median / det_total if abs(det_total) > 1e-12 else 1.0

    port = port_raw * factor
    cat_pp = dec["cat_pp"] * factor
    reg_pp = dec["reg_pp"] * factor
    A_cell = dec["A_cell"] * factor
    w_cell = dec["W"].reshape(-1)

    bands, band_total = band_analysis(config, drivers, dec, factor)

    term = port[:, y_last]
    gross_hw = float(term[term < 0].sum())
    gross_tw = float(term[term > 0].sum())
    order_abs = np.argsort(-np.abs(term))
    rank_abs = {int(j): i + 1 for i, j in enumerate(order_abs)}
    hw_ids = [int(j) for j in np.argsort(term) if term[j] < 0]
    tw_ids = [int(j) for j in np.argsort(-term) if term[j] > 0]
    rank_dir = {j: i + 1 for i, j in enumerate(hw_ids)}
    rank_dir.update({j: i + 1 for i, j in enumerate(tw_ids)})
    force_tot = {f: float(term[dec["fidx"] == fi].sum()) for fi, f in enumerate(forces)}
    force_hw = {f: float(term[(dec["fidx"] == fi) & (term < 0)].sum()) for fi, f in enumerate(forces)}
    force_tw = {f: float(term[(dec["fidx"] == fi) & (term > 0)].sum()) for fi, f in enumerate(forces)}
    force_rank = {}
    for fi in range(len(forces)):
        ids = [int(j) for j in np.argsort(-np.abs(term)) if dec["fidx"][j] == fi]
        force_rank.update({j: i + 1 for i, j in enumerate(ids)})
    cat_tot = cat_pp[:, :, y_last].sum(axis=0)                # (nc,) category shift at the means, scaled
    reg_tot = reg_pp[:, :, y_last].sum(axis=0)

    per_driver = {}
    for j, drv in enumerate(drivers):
        f = int(dec["fidx"][j])
        cells_term = A_cell[j, :, y_last] * w_cell                 # portfolio pp by cell
        top = [int(i) for i in np.argsort(-np.abs(cells_term))[:TOP_CELLS] if abs(cells_term[i]) > 0]
        own = float(term[j])
        cells = []
        for i in top:
            c, r = divmod(i, len(REGIONS))
            partners = []
            same_force = [k for k in range(len(drivers)) if dec["fidx"][k] == f and k != j and abs(A_cell[k, i, y_last]) > 0]
            same_force.sort(key=lambda k: -abs(A_cell[k, i, y_last]))
            for k in same_force[:TOP_PARTNERS]:
                partners.append({"id": drivers[k].id, "code": display_code(drivers[k].id, drivers[k].force),
                                 "pp": r4(A_cell[k, i, y_last] * w_cell[i] * 100)})
            cells.append({
                "category": cats[c], "region": REGIONS[r],
                "pp": r4(cells_term[i] * 100),
                "share_of_driver": r4(cells_term[i] / own) if own else 0.0,
                "cell_share_of_pool": r4(w_cell[i]),
                "cell_shift_pp": r4(dec["cell_shift"][i, y_last] * factor * 100),
                # the driver's share of the cell's own shift, both at the means and
                # unscaled (the reconciliation factor cancels in the ratio)
                "share_of_cell": r4(dec["A_cell"][j, i, y_last] / dec["cell_shift"][i, y_last]) if dec["cell_shift"][i, y_last] else 0.0,
                "n_eff": r4(dec["n_eff"][f, i]),
                "dampening": r4(dec["damp"][f, i]),
                "same_force_partners": partners,
            })
        # Over ALL 52 cells (the list above holds the six largest by points):
        # the cell where this driver is the largest share of the cell's shift,
        # so a reading can say "its largest share of any cell" truthfully.
        shares_all = np.array([
            (dec["A_cell"][j, i, y_last] / dec["cell_shift"][i, y_last]) if dec["cell_shift"][i, y_last] and abs(A_cell[j, i, y_last]) > 0 else 0.0
            for i in range(len(w_cell))])
        i_max = int(np.argmax(np.abs(shares_all)))
        c_max, r_max = divmod(i_max, len(REGIONS))
        max_share_cell = {
            "category": cats[c_max], "region": REGIONS[r_max],
            "share_of_cell": r4(shares_all[i_max]),
            "pp": r4(cells_term[i_max] * 100),
            "cell_shift_pp": r4(dec["cell_shift"][i_max, y_last] * factor * 100),
        } if abs(shares_all[i_max]) > 0 else None
        reach = float((w_cell * dec["E"][j, :]).sum())
        eff_damp = float((w_cell * dec["E"][j, :] * dec["damp"][f, :]).sum() / reach) if reach > 0 else 1.0
        # The additive term factorises into a year-independent constant and the
        # expected materialisation, so the share of the terminal-year effect
        # present in year y is exactly E[mat](y) / E[mat](terminal): monotone
        # and at most 1 (the portfolio path above still carries the small
        # year-specific cross-term spread, which is why it is not used here).
        e_last = float(dec["e_mat"][j, y_last])
        share_realised = {}
        for yi, y in enumerate(years):
            share_realised[str(y)] = r4(dec["e_mat"][j, yi] / e_last) if e_last > 0 else 0.0
        half_year = next((y for yi, y in enumerate(years) if e_last > 0 and dec["e_mat"][j, yi] / e_last >= 0.5), None)
        per_driver[drv.id] = {
            "code": display_code(drv.id, drv.force), "name": drv.name, "force": drv.force, "direction": drv.direction,
            "portfolio_pp": {str(y): r4(port[j, yi] * 100) for yi, y in enumerate(years)},
            "share_of_gross": r4(own / gross_hw) if own < 0 else (r4(own / gross_tw) if own > 0 else 0.0),
            "share_of_net": r4(own / median) if median else 0.0,
            "rank_abs": rank_abs[j], "rank_in_direction": rank_dir.get(j),
            # share of the force's gross headwind (tailwind) among the drivers of the same sign
            "force_share": (r4(own / force_hw[drv.force]) if own < 0 and force_hw[drv.force]
                            else (r4(own / force_tw[drv.force]) if own > 0 and force_tw[drv.force] else 0.0)),
            "force_rank": force_rank[j],
            "category_pp": {cats[c]: r4(cat_pp[j, c, y_last] * 100) for c in range(len(cats))},
            "category_share": {cats[c]: (r4(cat_pp[j, c, y_last] / cat_tot[c]) if cat_tot[c] else 0.0) for c in range(len(cats))},
            "region_pp": {REGIONS[r]: r4(reg_pp[j, r, y_last] * 100) for r in range(len(REGIONS))},
            "region_split": {REGIONS[r]: (r4(reg_pp[j, r, y_last] / own) if own else 0.0) for r in range(len(REGIONS))},
            "cells": cells,
            "max_share_cell": max_share_cell,
            "timing": {
                "expected_materialization": {str(y): r4(dec["e_mat"][j, yi]) for yi, y in enumerate(years)},
                "share_realised": share_realised,
                "half_year": half_year,
            },
            "band": bands[j],
            # the timing scores the reading may name through placeholders, so
            # that a re-scored peak year cannot leave a typed year in a text
            "scoring": {"peak_year": drv.peak_year, "diffusion_curve": drv.diffusion_curve,
                        "start_year": drv.start_year},
            "structure": {
                "pass_through": r4(dec["fw"][f] * dec["att"][f]),
                "reach": r4(reach),
                "effective_dampening": r4(eff_damp),
                "raw_shift": r4(float(dec["s_mean"][j])),
            },
        }

    record = {
        "generated_by": "scripts/build_driver_analysis.py",
        "record_version": 1,
        "built": _dt.date.today().isoformat(),
        "method": (
            "Expected-value decomposition of the engine's compounding arithmetic at the score "
            "means (Beta mean = probability/6, materialisation averaged over each driver's peak "
            "jitter), cross-terms spread in proportion to the additive terms, rolled up with the "
            "cell gross-profit shares. The sum reproduces the reference run's terminal-year "
            "mean; every contribution is then scaled by one factor, the run's median over that "
            "mean (the skew of the portfolio distribution), so the contributions add up to the "
            "median headline the dashboard shows. Attribution, not simulation."
        ),
        "basis": {
            "drivers": len(drivers),
            "model_version": reference["model_version"],
            "cell_weights_source": config.cell_weights_source,
            "attenuation_source": config.attenuation_source,
            "force_weights": {f: r4(config.force_weights[f]) for f in forces},
            "per_force_attenuation": {f: r4(config.per_force_attenuation[f]) for f in forces},
            "within_force_overlap": {f: r4(config.within_force_overlap.get(f, 0.0)) for f in forces},
            "regionless_drivers": dec["regionless"],
        },
        "reference_run": reference,
        # the production run the reference reproduces (from its QA workbook,
        # --workbook FILE); None until one is recorded
        "persisted_run": persisted,
        "reconciliation": {
            "terminal_year": TERMINAL_YEAR,
            "deterministic_total_pp": r4(det_total * 100),
            "reference_median_pp": r4(median * 100),
            "factor": round(factor, 6),
        },
        "totals": {
            "net_pp": r4(median * 100),
            "gross_headwind_pp": r4(gross_hw * 100),
            "gross_tailwind_pp": r4(gross_tw * 100),
            "headwinds": len(hw_ids), "tailwinds": len(tw_ids),
            "by_force_pp": {f: r4(force_tot[f] * 100) for f in forces},
            "by_force_headwind_pp": {f: r4(force_hw[f] * 100) for f in forces},
            "by_force_tailwind_pp": {f: r4(force_tw[f] * 100) for f in forces},
            "by_category_pp": {cats[c]: r4(cat_tot[c] * 100) for c in range(len(cats))},
            "by_region_pp": {REGIONS[r]: r4(reg_tot[r] * 100) for r in range(len(REGIONS))},
            # the mix the roll-up used, so a reading can name a region's share of the pool
            "pool_region_shares": {REGIONS[r]: r4(dec["W"][:, r].sum()) for r in range(len(REGIONS))},
            "pool_category_shares": {cats[c]: r4(dec["W"][c, :].sum()) for c in range(len(cats))},
            "band_p10_p90_pp": [r4(reference["portfolio_terminal"]["p10"] * 100),
                                r4(reference["portfolio_terminal"]["p90"] * 100)],
            # the additive-sum band of the band analysis, next to the run's band above
            "band_additive_sum": band_total,
        },
        "drivers": per_driver,
    }
    return record


# ── parity fixture for the TypeScript port ─────────────────────────────

def build_fixture(config: ModelConfig, drivers: list, dec: dict) -> dict:
    """Inputs exactly as the frontend receives them and the unscaled
    decomposition (no reconciliation factor), full float precision."""
    cats, years = dec["cats"], dec["years"]
    W = dec["W"]
    inputs = {
        "drivers": [{
            "id": drv.id, "force": drv.force,
            "normalized_score": float(drv.normalized_score),
            "category_exposure": {k: int(v) for k, v in drv.category_exposure.items()},
            "regional_exposure": {k: int(v) for k, v in (drv.regional_exposure or {}).items()},
            "uncertainty": drv.uncertainty, "peak_year": drv.peak_year,
            "diffusion_curve": drv.diffusion_curve, "start_year": drv.start_year,
        } for drv in drivers],
        "categories": list(cats),
        "config": {
            "force_weights": {f: float(config.force_weights[f]) for f in FORCES},
            "per_force_attenuation": {f: float(config.per_force_attenuation[f]) for f in FORCES},
            "within_force_overlap": {f: float(config.within_force_overlap.get(f, 0.0)) for f in FORCES},
            "base_year": int(config.base_year), "path_years": list(years),
            "peak_year_jitter": int(config.peak_year_jitter),
        },
        "cellWeights": {cat: {r: float(W[c, ri]) for ri, r in enumerate(REGIONS)} for c, cat in enumerate(cats)},
    }
    expected = {
        "byDriver": {
            drv.id: {
                "portfolio": {str(y): float(dec["port"][j, yi]) for yi, y in enumerate(years)},
                "category": {cat: {str(y): float(dec["cat_pp"][j, c, yi]) for yi, y in enumerate(years)}
                             for c, cat in enumerate(cats)},
            } for j, drv in enumerate(drivers)
        },
        "portfolioShift": {str(y): float(dec["port"][:, yi].sum()) for yi, y in enumerate(years)},
        "categoryShift": {cat: {str(y): float(dec["cat_pp"][:, c, yi].sum()) for yi, y in enumerate(years)}
                          for c, cat in enumerate(cats)},
    }
    return {"generated_by": "scripts/build_driver_analysis.py", "inputs": inputs, "expected": expected}


# ── main ───────────────────────────────────────────────────────────────

def _load_json(path: str) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _dump_json(path: str, obj: dict) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, indent=1, ensure_ascii=False)
        fh.write("\n")


def _strip_volatile(rec: dict) -> dict:
    r = json.loads(json.dumps(rec))
    r.pop("built", None)
    r.get("reference_run", {}).pop("run_at_utc", None)
    r.get("reference_run", {}).pop("numerics_backend", None)
    return r


# Tolerances for comparing a committed artefact with its recomputation. The
# arithmetic is deterministic, but libm and BLAS differ by an ulp or so between
# platforms (an arm64 laptop, CI's x86_64), and a value that sits on
# a rounding boundary (a triangular average of rational curve values can land
# exactly on a 5 in the fifth decimal) then rounds differently. Every field of
# the record is rounded to four decimals or finer, so one unit in the fourth
# decimal admits exactly such flips and nothing a reader could notice; the
# rendered texts are still compared verbatim, and the parity fixture, which
# carries full doubles, is held to the same relative tolerance as the vitest
# replay of it (tests/frontend/driverAttribution.test.ts).
RECORD_ATOL = 1.01e-4
FIXTURE_ATOL, FIXTURE_RTOL = 1e-12, 1e-9


def records_match(a, b, atol: float, rtol: float = 0.0) -> bool:
    """Structural equality with a numeric tolerance on leaf numbers."""
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(records_match(a[k], b[k], atol, rtol) for k in a)
    if isinstance(a, list) and isinstance(b, list):
        return len(a) == len(b) and all(records_match(x, y, atol, rtol) for x, y in zip(a, b))
    if isinstance(a, bool) or isinstance(b, bool):
        return isinstance(a, bool) and isinstance(b, bool) and a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) <= atol + rtol * abs(float(b))
    return a == b


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true", help="verify the committed record and texts, write nothing")
    ap.add_argument("--reference", action="store_true", help="rerun the reference run (the production recipe)")
    ap.add_argument("--chains", type=int, default=PRODUCTION_CHAINS)
    ap.add_argument("--iterations", type=int, default=PRODUCTION_ITERATIONS)
    ap.add_argument("--workbook", default=None,
                    help="the QA workbook of the production run the reference should reproduce")
    args = ap.parse_args(argv)

    drivers = load_drivers()
    config = ModelConfig()

    committed = _load_json(ATTRIBUTION_PATH) if os.path.exists(ATTRIBUTION_PATH) else None
    current_hash = inputs_hash(config, drivers)
    if args.reference or committed is None:
        print(f"reference run: {args.chains} x {args.iterations} iterations, seed {PRODUCTION_SEED} ...", flush=True)
        reference = reference_run(config, drivers, args.chains, args.iterations)
        print(f"  terminal-year portfolio median {reference['portfolio_terminal']['median']*100:+.3f} pp "
              f"[{reference['portfolio_terminal']['p10']*100:+.3f}, {reference['portfolio_terminal']['p90']*100:+.3f}]")
    else:
        reference = committed["reference_run"]
        if reference.get("inputs_hash") != current_hash and not args.check:
            print("REFUSING: the drivers' scoring state or the configuration changed since the recorded reference "
                  "run, so its median is not the current model's. Re-run with --reference --workbook FILE after "
                  "the production run.")
            return 1

    persisted = read_workbook(args.workbook) if args.workbook else (committed or {}).get("persisted_run")
    record = build_record(config, drivers, reference, persisted)
    fixture = build_fixture(config, drivers, decompose(config, drivers))
    analysis = _load_json(ANALYSIS_PATH) if os.path.exists(ANALYSIS_PATH) else {"drivers": {}}
    texts = render_all(record, analysis) if analysis.get("drivers") else {}

    if args.check:
        problems = []
        if committed is None:
            problems.append(f"{os.path.relpath(ATTRIBUTION_PATH, REPO)} is missing")
        elif not records_match(_strip_volatile(committed), _strip_volatile(record), RECORD_ATOL):
            problems.append(f"{os.path.relpath(ATTRIBUTION_PATH, REPO)} differs from the recomputed record "
                            f"(regenerate: python3 scripts/build_driver_analysis.py)")
        if committed is not None and committed["reference_run"].get("inputs_hash") != current_hash:
            problems.append("the drivers' scoring state or the configuration changed since the committed reference "
                            "run (its inputs digest differs): re-run with --reference --workbook FILE")
        if persisted:
            for line in reference_reproduces(reference, persisted):
                problems.append(f"the reference run does not reproduce production run {persisted.get('id')}: {line}")
        for code, entry in analysis.get("drivers", {}).items():
            driver_id = next((i for i, d in record["drivers"].items() if d["code"] == code), None)
            if driver_id and entry.get("text") != texts.get(driver_id):
                problems.append(f"{code}: rendered text differs from the committed one (regenerate)")
        if not os.path.exists(FIXTURE_PATH):
            problems.append(f"{os.path.relpath(FIXTURE_PATH, REPO)} is missing")
        elif not records_match(_load_json(FIXTURE_PATH), json.loads(json.dumps(fixture)), FIXTURE_ATOL, FIXTURE_RTOL):
            problems.append(f"{os.path.relpath(FIXTURE_PATH, REPO)} differs from the recomputed fixture (regenerate)")
        if problems:
            print("driver analysis check FAILED:")
            for p in problems:
                print("  -", p)
            return 1
        print("driver analysis check OK: record, fixture and texts match the driver base and the configuration.")
        return 0

    if persisted:
        mismatch = reference_reproduces(reference, persisted)
        if mismatch:
            print(f"WARNING: the reference run does not reproduce production run {persisted.get('id')} "
                  f"({len(mismatch)} quantile(s) differ, first: {mismatch[0]}). Either the drivers or the "
                  f"configuration changed since that run, or the workbook is not this run's: record the "
                  f"new production run with --workbook, or run run_production.py first.")
        else:
            print(f"reference run reproduces persisted run {persisted.get('id')} ({persisted.get('generated')}) "
                  f"at every terminal-year category quantile")
    _dump_json(ATTRIBUTION_PATH, record)
    print(f"wrote {os.path.relpath(ATTRIBUTION_PATH, REPO)} "
          f"({len(record['drivers'])} drivers, factor {record['reconciliation']['factor']:.4f})")
    _dump_json(FIXTURE_PATH, fixture)
    print(f"wrote {os.path.relpath(FIXTURE_PATH, REPO)} (parity fixture for lib/driverAttribution.ts)")
    if texts:
        for code, entry in analysis["drivers"].items():
            driver_id = next(i for i, d in record["drivers"].items() if d["code"] == code)
            entry["text"] = texts[driver_id]
        _dump_json(ANALYSIS_PATH, analysis)
        print(f"wrote {os.path.relpath(ANALYSIS_PATH, REPO)} ({len(texts)} rendered texts)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
