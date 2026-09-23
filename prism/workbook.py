"""The QA workbook written with every production run.

scripts/run_production.py writes one workbook per run so the numbers can be
checked offline, next to the run the dashboard serves. It holds percentage
shifts and gross-profit shares only, never an absolute financial figure.

Sheets: Shift Matrix (category paths with percentiles), Regional Shift
(median per category and region), Cell Weights (the roll-up shares and, for
the estimated mix, its graded derivation), Velocity, and Metadata (the run's
identity and the notes a reader needs to interpret the numbers).
read_workbook reads the run back from the Shift Matrix and Metadata sheets
for scripts/build_driver_analysis.py, so those two layouts are a contract
between the two functions below.
"""
import json
import logging
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from prism.config import (
    CATEGORIES,
    FORCES,
    REGIONS,
    UNCERTAINTY_KAPPA,
    UNCERTAINTY_PEAK_JITTER,
    ModelConfig,
)

logger = logging.getLogger(__name__)

ESTIMATE_RECORD = Path(__file__).resolve().parents[1] / "data" / "calibration" / "cell_weights_estimate.json"

PERCENTILES = ("p10", "p25", "median", "p75", "p90")

# The dashboard's colours (lib/theme.ts) as solid fills, since a workbook has
# no transparency: the navy ink, the surface tint, and the shift pair at 12
# percent over white.
INK, ON_INK, SURFACE_LOW = "00345E", "FFFFFF", "EFF4FF"
EXPANSION_TINT, CONTRACTION_TINT, RULE = "E4EFE8", "F3E8E8", "E6EBEF"

HEADER_FILL = PatternFill(start_color=INK, end_color=INK, fill_type="solid")
HEADER_FONT = Font(name="Inter", color=ON_INK, bold=True, size=10)
DATA_FONT = Font(name="Inter", color=INK, size=10)
LABEL_FONT = Font(name="Inter", size=9)
BOLD_FONT = Font(name="Inter", size=9, bold=True)
EXPANSION_FILL = PatternFill(start_color=EXPANSION_TINT, end_color=EXPANSION_TINT, fill_type="solid")
CONTRACTION_FILL = PatternFill(start_color=CONTRACTION_TINT, end_color=CONTRACTION_TINT, fill_type="solid")
NEUTRAL_FILL = PatternFill(start_color=SURFACE_LOW, end_color=SURFACE_LOW, fill_type="solid")
BORDER = Border(*(Side(style="thin", color=RULE) for _ in range(4)))

#: Notes printed on the Metadata sheet, in this order.
READING_NOTES = (
    ("Reading note",
     "PRISM holds strategy constant: the simulation propagates external drivers only and leaves out "
     "management response (price increases, innovation launches, mix shifts, competitive reaction). "
     "A negative total is a headwind to today's business if nothing changes, not a forecast that the "
     "pool will shrink."),
    ("Scale",
     "Per force, the pass-through is the force weight (1/6) times its attenuation (about 0.4 to 0.5), "
     "roughly 7 percent. A near-certain driver touching 20 percent of a category's GP1 moves that "
     "category by about 1.4 percent at full materialisation. Read the cells as a conservative, "
     "comparable index and apply them as GP1 projected = GP1 x (1 + shift)."),
    ("Regional shift",
     "The shift is solved per category and region: a driver reaches a (category, region) cell weighted "
     "by its category exposure times its regional exposure, so a regionally concentrated driver moves "
     "only its regions' slice of the pool."),
    ("Cell weights",
     "Category and portfolio numbers roll the 52 category x region cells up with each cell's share of "
     "HCB gross profit (sheet Cell Weights): a category is the share-weighted average of its regional "
     "cells, the portfolio the share-weighted sum of all cells."),
    ("Value-chain attribution",
     "Each driver's contribution is assigned wholly to the value-chain stage experts located as its "
     "epicentre. Propagation along the chain is not modelled."),
    ("Uncertainty",
     "P10 to P90 bands are Monte Carlo percentiles. The concentration of a driver's Beta prior follows "
     f"its uncertainty score 0 to 5 ({', '.join(str(UNCERTAINTY_KAPPA[u]) for u in range(6))}), and so "
     f"does its peak-year jitter ({', '.join(str(UNCERTAINTY_PEAK_JITTER[u]) for u in range(6))} years). "
     "A driver's start year gates its onset."),
    ("Score to prior",
     "Probability scores 1 to 5 give Beta prior means of 0.17, 0.33, 0.50, 0.67 and 0.83 (score / 6): "
     "a deliberate shrinkage, so a 5 means five in six, not certainty."),
    ("Copula",
     "Correlations are set on the latent Gaussian scale; the realised correlation of the probability "
     "draws is slightly lower (about 0.28 at 0.30)."),
    ("Confidentiality",
     "This file holds percentage shifts and gross-profit shares only; no absolute financial figure "
     "appears in it."),
)


def write_workbook(output_path: str, result: dict, config: ModelConfig, *,
                   run_id: Optional[int] = None, driver_count: Optional[int] = None) -> None:
    """Write the QA workbook for an engine result (``run_multichain`` output)."""
    wb = openpyxl.Workbook()
    _shift_matrix(wb.active, result, config)
    _regional_shift(wb.create_sheet("Regional Shift"), result, config)
    _cell_weights(wb.create_sheet("Cell Weights"), result, config)
    _velocity(wb.create_sheet("Velocity"), result, config)
    _metadata(wb.create_sheet("Metadata"), result, config, run_id, driver_count)
    wb.save(output_path)
    logger.debug("QA workbook written to %s", output_path)


# ── Helpers ─────────────────────────────────────────────────────────────────

def _header(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _shift_cell(ws, row: int, col: int, value: float) -> None:
    """A percentage cell coloured by direction (beyond +/-0.5 percent)."""
    cell = ws.cell(row=row, column=col, value=round(float(value), 6))
    cell.number_format = "0.00%"
    cell.font = DATA_FONT
    cell.alignment = Alignment(horizontal="center")
    cell.border = BORDER
    cell.fill = EXPANSION_FILL if value > 0.005 else CONTRACTION_FILL if value < -0.005 else NEUTRAL_FILL


@lru_cache(maxsize=1)
def _estimate_record() -> dict:
    return json.loads(ESTIMATE_RECORD.read_text(encoding="utf-8"))


# ── Sheets ──────────────────────────────────────────────────────────────────

def _shift_matrix(ws, result: dict, config: ModelConfig) -> None:
    """Two header rows (year, then percentile) and one row per category."""
    ws.title = "Shift Matrix"
    ws.merge_cells("A1:A2")
    _header(ws, 1, 1, "Category")
    col = 2
    for year in config.path_years:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(PERCENTILES) - 1)
        _header(ws, 1, col, str(year))
        for p in PERCENTILES:
            _header(ws, 2, col, p)
            col += 1
    for row, cat in enumerate(CATEGORIES, 3):
        ws.cell(row=row, column=1, value=cat).font = Font(name="Inter", bold=True, size=10)
        path = (result["shift_matrix"].get(cat) or {}).get("path", {})
        col = 2
        for year in config.path_years:
            q = path.get(year, {})
            for p in PERCENTILES:
                _shift_cell(ws, row, col, q.get(p, 0.0))
                col += 1
    ws.column_dimensions["A"].width = 18
    for c in range(2, col):
        ws.column_dimensions[get_column_letter(c)].width = 12


def _regional_shift(ws, result: dict, config: ModelConfig) -> None:
    """Median shift per (category, region) and year, labelled with the cell's share."""
    _header(ws, 1, 1, "Category")
    _header(ws, 1, 2, "Region (share of pool)")
    for j, year in enumerate(config.path_years, 3):
        _header(ws, 1, j, str(year))
    weights = result["cell_weights_used"]
    row = 2
    for cat in CATEGORIES:
        for region in REGIONS:
            share = (weights.get(cat) or {}).get(region, 0.0)
            ws.cell(row=row, column=1, value=cat).font = LABEL_FONT
            ws.cell(row=row, column=2, value=f"{region} ({share * 100:.1f}%)").font = LABEL_FONT
            path = ((result["regional_shift_matrix"].get(cat) or {}).get(region) or {}).get("path", {})
            for j, year in enumerate(config.path_years, 3):
                _shift_cell(ws, row, j, path.get(year, {}).get("median", 0.0))
            row += 1
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 24
    for c in range(3, 3 + len(config.path_years)):
        ws.column_dimensions[get_column_letter(c)].width = 11


def _cell_weights(ws, result: dict, config: ModelConfig) -> None:
    """The share matrix the run applied, its marginals and its source."""
    weights = result["cell_weights_used"]
    _header(ws, 1, 1, "Category")
    for j, region in enumerate(REGIONS, 2):
        _header(ws, 1, j, region)
    _header(ws, 1, len(REGIONS) + 2, "Category share")
    row = 2
    for cat in CATEGORIES:
        shares = weights.get(cat) or {}
        ws.cell(row=row, column=1, value=cat).font = LABEL_FONT
        for j, region in enumerate(REGIONS, 2):
            ws.cell(row=row, column=j, value=round(float(shares.get(region, 0.0)), 6)).number_format = "0.00%"
        ws.cell(row=row, column=len(REGIONS) + 2,
                value=round(sum(float(v) for v in shares.values()), 6)).number_format = "0.00%"
        row += 1
    ws.cell(row=row, column=1, value="Region share").font = BOLD_FONT
    for j, region in enumerate(REGIONS, 2):
        total = sum(float((weights.get(c) or {}).get(region, 0.0)) for c in CATEGORIES)
        ws.cell(row=row, column=j, value=round(total, 6)).number_format = "0.00%"
    ws.cell(row=row + 2, column=1, value="Source: " + str(result.get("cell_weights_source") or ""))
    row = _estimate_derivation(ws, weights, config, row + 4)
    ws.column_dimensions["A"].width = 18
    for j in range(2, len(REGIONS) + 3):
        ws.column_dimensions[get_column_letter(j)].width = 14


def _estimate_derivation(ws, weights: dict, config: ModelConfig, row: int) -> int:
    """The graded derivation of the estimated mix, written only when the run used it.

    A reader holding the file offline must be able to see that these shares
    are an estimate and how each one was reached; a source line alone invites
    them to be read as P&L. A run on the finance figures gets no such block.
    """
    record = _estimate_record()
    cells = record["cells"]
    if set(weights) != set(cells) or any(
            abs(float((weights.get(c) or {}).get(r, 0.0)) - float(cells[c][r])) >= 1e-9
            for c in cells for r in cells[c]):
        return row
    inputs = record["inputs"]

    ws.cell(row=row, column=1, value="How these shares were built").font = BOLD_FONT
    row += 1
    for line in (
        "share = (category's share of its block x the block's share of HCB) x GP1 margin index x "
        "that category's own regional mix, normalised to 1.",
        "This is an ESTIMATE from public reporting and category knowledge. It is NOT Henkel P&L and "
        "must not be quoted as such.",
        "Each category has its own regional mix, so the matrix is not separable: regional "
        "concentration moves the category numbers.",
    ):
        ws.cell(row=row, column=1, value=line).font = LABEL_FONT
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="Grade").font = BOLD_FONT
    ws.cell(row=row, column=2, value="Meaning").font = BOLD_FONT
    row += 1
    for grade, meaning in record["grade_legend"].items():
        ws.cell(row=row, column=1, value=grade).font = LABEL_FONT
        ws.cell(row=row, column=2, value=meaning).font = LABEL_FONT
        row += 1
    row += 1

    split = inputs["block_split"]
    ws.cell(row=row, column=1, value="Block split").font = BOLD_FONT
    ws.cell(row=row, column=2, value=f"LHC {float(split['LHC']) * 100:.0f}% / "
                                     f"Hair & Body {float(split['Hair & Body']) * 100:.0f}%").font = LABEL_FONT
    ws.cell(row=row, column=3, value=str(split["grade"])).font = LABEL_FONT
    ws.cell(row=row, column=4, value=str(split["note"])).font = LABEL_FONT
    row += 2

    for j, head in enumerate(("Category", "Share of block", "Grade", "Margin index", "Grade", "Basis"), 1):
        _header(ws, row, j, head)
    row += 1
    for cat in CATEGORIES:
        mix = inputs["category_mix"][cat]
        margin = inputs["margin_index"][cat]
        ws.cell(row=row, column=1, value=cat).font = LABEL_FONT
        ws.cell(row=row, column=2, value=float(mix["share_of_block"])).number_format = "0.0%"
        ws.cell(row=row, column=3, value=str(mix["grade"])).font = LABEL_FONT
        ws.cell(row=row, column=4, value=float(margin["index"])).number_format = "0.00"
        ws.cell(row=row, column=5, value=str(margin["grade"])).font = LABEL_FONT
        ws.cell(row=row, column=6, value=str(mix["note"])).font = LABEL_FONT
        row += 1
    row += 1

    for j, head in enumerate(["Category", *REGIONS, "Grade", "Basis"], 1):
        _header(ws, row, j, head)
    row += 1
    for cat in CATEGORIES:
        entry = inputs["region_mix"][cat]
        ws.cell(row=row, column=1, value=cat).font = LABEL_FONT
        for j, region in enumerate(REGIONS, 2):
            ws.cell(row=row, column=j, value=float(entry["mix"][region])).number_format = "0%"
        ws.cell(row=row, column=len(REGIONS) + 2, value=str(entry["grade"])).font = LABEL_FONT
        ws.cell(row=row, column=len(REGIONS) + 3, value=str(entry["note"])).font = LABEL_FONT
        row += 1
    for note in (inputs.get("margin_index_note"), inputs.get("region_mix_note")):
        if note:
            row += 1
            ws.cell(row=row, column=1, value=str(note)).font = LABEL_FONT
    row += 2
    ws.cell(row=row, column=1, value="Replace all of this with the finance figures: "
                                     "run_production.py --cell-weights FILE.").font = BOLD_FONT
    return row + 1


def _velocity(ws, result: dict, config: ModelConfig) -> None:
    """Median year-on-year change per category, with a coarse shape label."""
    years = config.path_years[1:]
    for j, head in enumerate(["Category", "Path shape", *(f"Velocity {y}" for y in years)], 1):
        _header(ws, 1, j, head)
    for row, cat in enumerate(CATEGORIES, 2):
        velocity = (result["shift_matrix"].get(cat) or {}).get("velocity", {})
        values = [float(velocity.get(y, {}).get("median", 0.0)) for y in years]
        if all(abs(v) < 0.001 for v in values):
            shape = "flat"
        elif abs(values[-1]) > abs(values[0]) * 1.5:
            shape = "accelerating"
        elif abs(values[0]) > abs(values[-1]) * 1.5:
            shape = "decelerating"
        else:
            shape = "gradual"
        ws.cell(row=row, column=1, value=cat).font = Font(name="Inter", bold=True)
        ws.cell(row=row, column=2, value=shape)
        for j, v in enumerate(values, 3):
            ws.cell(row=row, column=j, value=round(v, 6)).number_format = "0.00%"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 14


def _metadata(ws, result: dict, config: ModelConfig, run_id: Optional[int],
              driver_count: Optional[int]) -> None:
    """Label in column A, value in column B: the run's identity, then the reading notes."""
    attenuation = ", ".join(f"{f} {config.per_force_attenuation[f]:.3f}" for f in FORCES)
    regions = ", ".join(f"{r} {w * 100:.1f}%" for r, w in result["region_weights_used"].items())
    rows = [
        ("Run id", run_id),
        ("Generated", datetime.now(timezone.utc).isoformat(timespec="seconds")),
        ("Model version", result.get("model_version")),
        ("Engine", f"{result.get('engine_name')} (Beta priors, Gaussian copula)"),
        ("Numerics", result.get("numerics_backend")),
        ("Iterations", result.get("iterations")),
        ("Chains", result.get("n_chains")),
        ("Seed", result.get("seed")),
        ("Drivers", driver_count),
        ("Path years", ", ".join(str(y) for y in config.path_years)),
        ("Attenuation source", f"{config.attenuation_source} (structured-judgment overlap correction, "
                               "not fitted to outcomes)"),
        ("Per-force attenuation", attenuation),
        ("Cell weights", str(result.get("cell_weights_source") or "")),
        ("Region shares", regions),
        (None, None),
        *READING_NOTES,
    ]
    wrap = Alignment(wrap_text=True, vertical="top")
    for row, (label, value) in enumerate(rows, 1):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value).alignment = wrap
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100


# ── Reading a workbook back ─────────────────────────────────────────────────

def read_workbook(path: str, terminal_year: int) -> dict:
    """The run a QA workbook records: its identity and, for every category, the
    P10, median and P90 of ``terminal_year``.

    scripts/build_driver_analysis.py keeps this in the attribution record, so
    that its check proves the reference run reproduces the run the dashboard
    serves, and a reader of the record knows which run that is.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    meta = {str(row[0]).strip(): row[1] for row in wb["Metadata"].iter_rows(values_only=True)
            if row and row[0] is not None and len(row) > 1}
    rows = list(wb["Shift Matrix"].iter_rows(values_only=True))
    header, sub = rows[0], rows[1]
    start = next((i for i, v in enumerate(header) if str(v) == str(terminal_year)), None)
    if start is None or tuple(str(x) for x in sub[start:start + len(PERCENTILES)]) != PERCENTILES:
        raise ValueError(f"{path}: no {terminal_year} percentiles in the Shift Matrix sheet")
    categories = {}
    for r in rows[2:]:
        if r and r[0] is not None and r[start] is not None:
            q = dict(zip(PERCENTILES, r[start:start + len(PERCENTILES)]))
            categories[str(r[0])] = {k: round(float(q[k]), 6) for k in ("p10", "median", "p90")}
    return {
        "id": int(meta["Run id"]) if meta.get("Run id") is not None else None,
        "generated": str(meta.get("Generated") or ""),
        "model_version": str(meta.get("Model version") or ""),
        "pooled_iterations": int(meta["Iterations"]) if meta.get("Iterations") is not None else None,
        "workbook": Path(path).name,
        "categories_terminal": categories,
    }
