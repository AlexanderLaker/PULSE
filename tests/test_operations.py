"""The operator's tools: the production run, its QA workbook, the .env loader and the admin script."""
import json
import os
import sys
from pathlib import Path

import pytest

from support import database_awaiting_a_migration, load_module

REPO = Path(__file__).resolve().parents[1]


@pytest.fixture
def seeded(monkeypatch, tmp_path):
    """A fresh SQLite database holding the seed drivers."""
    monkeypatch.setenv("PRISM_DB_PATH", str(tmp_path / "prism.db"))
    import prism.database as dbm
    from prism.seed_drivers import get_seed_drivers
    dbm.init_db()
    dbm.insert_drivers(get_seed_drivers())
    return dbm


@pytest.fixture
def production(monkeypatch, tmp_path):
    module = load_module("scripts/run_production.py")
    monkeypatch.setattr(module, "RUNS_DIR", tmp_path / "runs")
    return module


class TestProductionRun:
    def test_refuses_without_a_database(self, production, monkeypatch):
        monkeypatch.setattr(production, "USE_POSTGRES", False)
        assert production.main([]) == 1

    def test_refuses_a_database_awaiting_a_migration(self, production, monkeypatch, tmp_path):
        awaiting = tmp_path / "awaiting.db"
        database_awaiting_a_migration(awaiting)
        monkeypatch.setenv("PRISM_DB_PATH", str(awaiting))
        assert production.main(["--allow-sqlite"]) == 1

    def test_an_empty_local_file_is_seeded_with_the_driver_base(self, production, monkeypatch, tmp_path):
        monkeypatch.setenv("PRISM_DB_PATH", str(tmp_path / "empty.db"))
        import prism.database as dbm
        dbm.init_db()
        assert production.main(["--allow-sqlite", "--iterations", "200"]) == 0
        assert len(dbm.load_drivers()) == 51
        assert dbm.load_latest_run()["results"]["meta"]["driver_count"] == 51

    def test_an_empty_postgresql_database_is_refused(self, production, monkeypatch, tmp_path):
        """PostgreSQL is never seeded by the run: an empty one is a new database the API has not started
        on, or a production database that lost its drivers and is restored from its backup."""
        monkeypatch.setenv("PRISM_DB_PATH", str(tmp_path / "empty.db"))
        monkeypatch.setattr(production, "USE_POSTGRES", True)
        import prism.database as dbm
        dbm.init_db()
        assert production.main(["--iterations", "200", "--allow-test-size"]) == 2
        assert dbm.load_drivers() == []

    def test_on_postgresql_only_the_production_size_runs(self, production, seeded, monkeypatch, caplog):
        """A run on PostgreSQL is served as soon as it is stored, so a test size needs --allow-test-size."""
        monkeypatch.setattr(production, "USE_POSTGRES", True)
        for size in (["--iterations", "2000"], ["--chains", "2"]):
            assert production.main(size) == 5
        assert "--allow-test-size" in caplog.text
        assert seeded.load_latest_run() is None
        assert production.main(["--iterations", "300", "--allow-test-size"]) == 0
        assert seeded.load_latest_run()["iterations"] == 300

    def test_the_run_records_where_its_code_came_from(self, production, seeded, monkeypatch, tmp_path):
        import re
        import openpyxl
        assert production.main(["--allow-sqlite", "--iterations", "200"]) == 0
        run = seeded.load_latest_run()
        source = run["results"]["meta"]["source"]
        assert set(source) == {"commit", "dirty"}
        if source["commit"] is not None:   # a git checkout
            assert re.fullmatch(r"[0-9a-f]{40}", source["commit"]) and isinstance(source["dirty"], bool)
        workbook = next((tmp_path / "runs").glob(f"prism_run_{run['id']}_*.xlsx"))
        labels = {row[0]: row[1] for row in openpyxl.load_workbook(workbook)["Metadata"].iter_rows(values_only=True)}
        assert labels["Code"] == production.describe_source(source)

        def no_git(*args, **kwargs):
            raise FileNotFoundError("git")
        monkeypatch.setattr(production.subprocess, "run", no_git)
        assert production.source_state() == {"commit": None, "dirty": None}
        assert production.describe_source(production.source_state()) == "not a git checkout"

    def test_a_run_is_persisted_with_the_full_contract(self, production, seeded, tmp_path):
        from prism.api.runs import META_KEYS, REQUIRED_KEYS, run_status
        assert production.main(["--allow-sqlite", "--iterations", "300", "--chains", "2"]) == 0
        run = seeded.load_latest_run()
        assert run_status(run)["reason"] == "ok"
        assert set(REQUIRED_KEYS) <= set(run["results"])
        meta = run["results"]["meta"]
        assert set(META_KEYS) <= set(meta) and meta["driver_count"] == 51 and meta["chains"] == 2
        assert set(meta["driver_fingerprint"]) == {d.id for d in seeded.load_drivers()}
        assert run["iterations"] == 300
        assert run["results"]["integrity_events"] == []   # the first run has nothing to compare with
        workbook = list((tmp_path / "runs").glob(f"prism_run_{run['id']}_*.xlsx"))
        assert len(workbook) == 1

        assert production.main(["--allow-sqlite", "--iterations", "300", "--chains", "2"]) == 0
        events = seeded.load_latest_run()["results"]["integrity_events"]
        assert [e["type"] for e in events] == ["input_drift"]
        assert f"vs run #{run['id']}" in events[0]["message"] and "no changes" in events[0]["message"]

    @pytest.mark.parametrize("size", [["--iterations", "50"], ["--iterations", "0"], ["--chains", "0"],
                                      ["--chains", "11"]])
    def test_a_run_size_out_of_range_is_exit_5(self, production, seeded, size, caplog):
        assert production.main(["--allow-sqlite", *size]) == 5
        assert "invalid run size" in caplog.text
        assert seeded.load_latest_run() is None

    def test_a_previous_run_that_cannot_be_compared_with_is_recorded(self, production, seeded):
        """A comparison that did not happen is a warning with the run, never the silence of a first run."""
        assert production.main(["--allow-sqlite", "--iterations", "200"]) == 0
        first = seeded.load_latest_run()
        results = first["results"]
        del results["meta"]["driver_fingerprint"]["consumer_r01"]["u"]
        with seeded.get_db_connection() as conn:
            conn.cursor().execute(f"UPDATE simulation_runs SET results = {seeded.placeholder()} WHERE id = "
                                  f"{seeded.placeholder()}", (json.dumps(results), first["id"]))
        assert production.main(["--allow-sqlite", "--iterations", "200"]) == 0
        [event] = seeded.load_latest_run()["results"]["integrity_events"]
        assert event["severity"] == "warning" and event["detail"]["compared"] is False
        assert "its driver fingerprint is incomplete for 1 driver" in event["message"]

    def test_a_driver_that_cannot_be_simulated_is_exit_6(self, production, seeded, caplog):
        """The database refuses a missing or off-scale score itself; what a change made outside PRISM can
        still leave behind (a missing exposure row, a peak year moved onto the onset) stops the run."""
        with seeded.get_db_connection() as conn:
            conn.cursor().execute("DELETE FROM driver_category_exposure "
                                  "WHERE driver_id = 'consumer_r02' AND category = 'LHC: IC'")
            conn.cursor().execute("UPDATE drivers SET peak_year = start_year WHERE id = 'technology_r01'")
        assert production.main(["--allow-sqlite", "--iterations", "200"]) == 6
        assert "driver consumer_r02: category_exposure has no score for LHC: IC" in caplog.text
        assert "driver technology_r01: peak year" in caplog.text
        assert seeded.load_latest_run() is None

    @pytest.mark.parametrize("lam_min", [-0.01, 0.0], ids=["negative", "zero"])
    def test_the_correlation_check_stops_the_run(self, production, seeded, monkeypatch, lam_min):
        """The copula needs a positive definite matrix: a smallest eigenvalue of zero stops the run too."""
        monkeypatch.setattr(production, "correlation_lambda_min", lambda *a: lam_min)
        assert production.main(["--allow-sqlite", "--iterations", "200"]) == 4
        assert seeded.load_latest_run() is None

    def test_a_failed_persist_is_exit_3(self, production, seeded, monkeypatch):
        def fail(*args, **kwargs):
            raise RuntimeError("database gone")
        monkeypatch.setattr(production, "save_simulation_run", fail)
        assert production.main(["--allow-sqlite", "--iterations", "200", "--chains", "1"]) == 3

    def test_a_bad_cell_weights_file_is_exit_5(self, production, seeded, tmp_path):
        bad = tmp_path / "weights.json"
        bad.write_text(json.dumps({"cells": {}}), encoding="utf-8")
        assert production.main(["--allow-sqlite", "--cell-weights", str(bad)]) == 5


class TestCellWeightsFile:
    def _file(self, tmp_path, cells, basis="gp1_share", source="Finance FY2025"):
        path = tmp_path / "weights.json"
        path.write_text(json.dumps({"source": source, "basis": basis, "cells": cells}), encoding="utf-8")
        return path

    def test_shares_and_absolute_figures(self, production, tmp_path):
        from prism.config import CATEGORIES, REGIONS
        equal = {c: {r: 1.0 / 52 for r in REGIONS} for c in CATEGORIES}
        shares, source = production.load_cell_weights_file(self._file(tmp_path, equal))
        assert source == "Finance FY2025" and shares["LHC: IC"]["Asia"] == pytest.approx(1.0 / 52)
        euros = {c: {r: 250.0 for r in REGIONS} for c in CATEGORIES}
        shares, source = production.load_cell_weights_file(self._file(tmp_path, euros, basis="gp1_absolute"))
        assert shares["LHC: IC"]["Asia"] == pytest.approx(1.0 / 52)
        assert "no amount is stored" in source

    @pytest.mark.parametrize("mutate", [
        lambda cells: cells.pop("LHC: IC"),
        lambda cells: cells["Hair: Care"].pop("Asia"),
        lambda cells: cells["Hair: Care"].update({"Mars": 0.1}),
        lambda cells: cells["Hair: Care"].update({"Asia": -1}),
        lambda cells: cells["Hair: Care"].update({"Asia": 0.9}),   # no longer sums to 1
    ])
    def test_defects_are_refused(self, production, tmp_path, mutate):
        from prism.config import CATEGORIES, REGIONS
        cells = {c: {r: 1.0 / 52 for r in REGIONS} for c in CATEGORIES}
        mutate(cells)
        with pytest.raises(ValueError):
            production.load_cell_weights_file(self._file(tmp_path, cells))

    def test_a_label_that_opens_like_the_estimate_is_refused(self, production, tmp_path):
        """The dashboard reads that opening as the built-in estimate."""
        from prism.config import CATEGORIES, REGIONS
        cells = {c: {r: 1.0 / 52 for r in REGIONS} for c in CATEGORIES}
        with pytest.raises(ValueError, match="built-in estimate"):
            production.load_cell_weights_file(self._file(tmp_path, cells, source="Estimate, not Henkel P&L (FY25)"))
        shares, source = production.load_cell_weights_file(self._file(tmp_path, cells, source="Estimated mix, Finance"))
        assert source == "Estimated mix, Finance"


class TestWorkbook:
    def _result(self, config, driver_set):
        from prism.engine import MonteCarloEngine
        return MonteCarloEngine(config, seed=42).run(driver_set, chains=2, iterations=200)

    def test_sheets_and_read_back(self, tmp_path, mock_model_config, mock_driver_set):
        import openpyxl
        from prism.workbook import read_workbook, write_workbook
        result = self._result(mock_model_config, mock_driver_set)
        path = tmp_path / "run.xlsx"
        write_workbook(str(path), result, mock_model_config, run_id=7, driver_count=5)
        wb = openpyxl.load_workbook(path)
        assert wb.sheetnames == ["Shift Matrix", "Regional Shift", "Cell Weights", "Velocity", "Metadata"]
        meta = {row[0]: row[1] for row in wb["Metadata"].iter_rows(values_only=True) if row[0]}
        assert meta["Run id"] == 7 and meta["Drivers"] == 5
        assert (meta["Iterations per chain"], meta["Chains"], meta["Pooled iterations"]) == (200, 2, 400)
        assert meta["Model version"] == result["model_version"]
        assert "Reading note" in meta and "Confidentiality" in meta
        # the notes quote the configuration the run used, never a typed figure
        from prism.config import FORCES
        pass_through = sorted(mock_model_config.force_weights[f] * mock_model_config.per_force_attenuation[f]
                              for f in FORCES)
        assert f"comes to {pass_through[0] * 100:.1f} to {pass_through[-1] * 100:.1f} percent" in meta["Scale"]
        assert f"at {mock_model_config.within_force_rho:.2f} within a force" in meta["Copula"]

        read = read_workbook(str(path), mock_model_config.path_years[-1])
        assert read["id"] == 7 and read["pooled_iterations"] == 400 and read["workbook"] == "run.xlsx"
        assert read["categories_terminal"]["LHC: FCN"]["median"] == pytest.approx(
            result["shift_matrix"]["LHC: FCN"]["path"][2030]["median"], abs=1e-6)
        with pytest.raises(ValueError, match="no 2035 percentiles"):
            read_workbook(str(path), 2035)

    def test_only_a_stored_runs_workbook_is_read_back(self, tmp_path, mock_model_config, mock_driver_set):
        """The attribution record names the run it reproduces, so a workbook
        without a run id (the run was not stored) is refused, not read as run None."""
        from prism.workbook import read_workbook, write_workbook
        path = tmp_path / "unsaved.xlsx"
        write_workbook(str(path), self._result(mock_model_config, mock_driver_set), mock_model_config)
        with pytest.raises(ValueError, match="the Metadata sheet has no Run id"):
            read_workbook(str(path), mock_model_config.path_years[-1])

    def test_the_estimate_carries_its_derivation(self, tmp_path, mock_model_config, mock_driver_set):
        import openpyxl
        from prism.workbook import write_workbook
        path = tmp_path / "estimate.xlsx"
        write_workbook(str(path), self._result(mock_model_config, mock_driver_set), mock_model_config)
        text = "\n".join(str(c) for row in openpyxl.load_workbook(path)["Cell Weights"].iter_rows(values_only=True)
                         for c in row if c is not None)
        assert "How these shares were built" in text and "NOT Henkel P&L" in text
        assert "Share of block" in text and "Margin index" in text
        assert "run_production.py --cell-weights FILE" in text

    def test_other_weights_get_no_derivation(self, tmp_path, mock_model_config, mock_driver_set):
        import openpyxl
        from prism.config import CATEGORIES, REGIONS
        from prism.workbook import write_workbook
        cfg = mock_model_config.copy_with(cell_weights={c: {r: 1.0 / 52 for r in REGIONS} for c in CATEGORIES},
                                          cell_weights_source="Finance FY2025")
        path = tmp_path / "finance.xlsx"
        write_workbook(str(path), self._result(cfg, mock_driver_set), cfg)
        text = "\n".join(str(c) for row in openpyxl.load_workbook(path)["Cell Weights"].iter_rows(values_only=True)
                         for c in row if c is not None)
        assert "How these shares were built" not in text and "Source: Finance FY2025" in text


class TestEnvFile:
    def test_the_file_is_loaded_and_the_shell_wins(self, tmp_path, monkeypatch):
        from prism import env_loader
        env = tmp_path / ".env"
        env.write_text("PRISM_TEST_FROM_FILE=file\nPRISM_TEST_BOTH=file\n", encoding="utf-8")
        monkeypatch.delenv("PRISM_TEST_FROM_FILE", raising=False)
        monkeypatch.setenv("PRISM_TEST_BOTH", "shell")
        try:
            assert env_loader.load_env_file(env) is True
            assert os.environ["PRISM_TEST_FROM_FILE"] == "file" and os.environ["PRISM_TEST_BOTH"] == "shell"
        finally:
            os.environ.pop("PRISM_TEST_FROM_FILE", None)

    def test_without_python_dotenv_the_file_is_skipped_with_a_warning(self, tmp_path, monkeypatch, caplog):
        """A host that installs only api/requirements.txt still starts the API."""
        from prism import env_loader
        env = tmp_path / ".env"
        env.write_text("PRISM_TEST_SKIPPED=1\n", encoding="utf-8")
        monkeypatch.setitem(sys.modules, "dotenv", None)   # importing dotenv now fails
        assert env_loader.load_env_file(env) is False
        assert "python-dotenv is not installed" in caplog.text and "PRISM_TEST_SKIPPED" not in os.environ

    def test_no_file_no_change(self, tmp_path):
        from prism import env_loader
        assert env_loader.load_env_file(tmp_path / ".env") is False


def test_the_admin_script_imports():
    module = load_module("scripts/promote_admin.py")
    assert callable(module.main)


def test_the_admin_script_requires_tls_unless_the_url_says_otherwise(monkeypatch, capsys):
    module = load_module("scripts/promote_admin.py")
    calls = []

    def connect(dsn, **kwargs):
        calls.append(kwargs)
        raise RuntimeError("not connecting in a test")
    monkeypatch.setattr(module.psycopg2, "connect", connect)
    for dsn in ("postgresql://user@db.example.org/prism", "postgresql://user@localhost/prism?sslmode=disable"):
        with pytest.raises(RuntimeError):
            module.main(["ana@example.com", "--dsn", dsn])
    assert calls == [{"sslmode": "require"}, {}]


def test_the_admin_script_names_a_failed_connection_by_its_category_only(capsys):
    """psycopg2's message about a malformed URL quotes part of it, password included."""
    module = load_module("scripts/promote_admin.py")
    assert module.main(["ana@example.com", "--dsn", "postgresql://owner:s3cr%zzet@localhost/prism"]) == 1
    err = capsys.readouterr().err
    assert "Cannot connect: the database URL is malformed" in err and "s3cr" not in err


@pytest.mark.skipif(not os.environ.get("PRISM_TEST_POSTGRES_URL"), reason="PRISM_TEST_POSTGRES_URL is not set")
def test_the_admin_script_changes_the_role_with_or_without_an_audit_log(capsys):
    """A database the API never started on has user_roles (the dashboard's) but no audit_log."""
    import uuid
    from urllib.parse import quote

    import psycopg2
    url = os.environ["PRISM_TEST_POSTGRES_URL"]
    schema = f"admin_{uuid.uuid4().hex[:8]}"
    admin = psycopg2.connect(url)
    admin.autocommit = True
    cur = admin.cursor()
    cur.execute(f"CREATE SCHEMA {schema}")
    try:
        cur.execute(f"CREATE TABLE {schema}.user_roles (clerk_user_id TEXT PRIMARY KEY, email TEXT, role TEXT, "
                    "updated_at TIMESTAMP)")
        cur.execute(f"INSERT INTO {schema}.user_roles VALUES ('u1', 'Ana@Example.com', 'viewer', NULL)")
        dsn = f"{url}{'&' if '?' in url else '?'}options={quote(f'-c search_path={schema}')}"
        module = load_module("scripts/promote_admin.py")
        assert module.main(["ana@example.com", "--dsn", dsn]) == 0
        assert "(no audit log yet, so not logged)" in capsys.readouterr().out
        cur.execute(f"SELECT role FROM {schema}.user_roles")
        assert cur.fetchone()[0] == "admin"

        cur.execute(f"CREATE TABLE {schema}.audit_log (id SERIAL PRIMARY KEY, action TEXT, entity_type TEXT, "
                    "entity_id TEXT, old_value TEXT, new_value TEXT, reason TEXT, user_id TEXT)")
        assert module.main(["ana@example.com", "--role", "viewer", "--dsn", dsn]) == 0
        cur.execute(f"SELECT action, old_value, new_value FROM {schema}.audit_log")
        assert cur.fetchall() == [("role_change", "admin", "viewer")]
        assert module.main(["nobody@example.com", "--dsn", dsn]) == 3
    finally:
        cur.execute(f"DROP SCHEMA {schema} CASCADE")
        admin.close()
